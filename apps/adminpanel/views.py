from django.shortcuts import render, redirect, get_object_or_404
from django.contrib import messages
from django.http import JsonResponse, HttpResponse
from django.views.decorators.http import require_POST
from django.utils import timezone
from .decorators import admin_required
from apps.users.models import User
from apps.properties.models import Property, PropertyImage
from apps.promotions.models import Promotion
from apps.dossiers.models import Dossier, EtapeGlobale, EtapeCochee, DossierHistorique, IntituleDossier
from apps.reservations.models import Reservation
from apps.chatbot.models import ChatbotKnowledge, ChatbotUnknown
import csv
import io
import datetime
import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

def _xl_val(cell):
    if cell is None:
        return ''
    val = cell.value
    if val is None:
        return ''
    return str(val).strip()


# ── DASHBOARD ──
@admin_required
def admin_dashboard(request):
    dossiers_alerte_list = Dossier.objects.filter(created_at__lte=timezone.now() - datetime.timedelta(days=45))
    dossiers_alerte = [d for d in dossiers_alerte_list if not d.is_complete()]
    ctx = {
        'total_clients': User.objects.filter(role='client').count(),
        'total_properties': Property.objects.count(),
        'total_dossiers': Dossier.objects.count(),
        'total_reservations': Reservation.objects.filter(status='pending').count(),
        'recent_reservations': Reservation.objects.order_by('-created_at')[:6],
        'recent_dossiers': Dossier.objects.select_related('client', 'intitule').order_by('-created_at')[:5],
        'unknown_questions': ChatbotUnknown.objects.count(),
        'alerte_45j_count': len(dossiers_alerte),
        'dossiers_alerte': dossiers_alerte[:3],
    }
    return render(request, 'adminpanel/dashboard.html', ctx)


# ── TERRAINS ──
@admin_required
def admin_properties(request):
    return render(request, 'adminpanel/properties/list.html', {'properties': Property.objects.prefetch_related('images').all()})


@admin_required
def admin_property_create(request):
    if request.method == 'POST':
        p = Property(title=request.POST.get('title'), description=request.POST.get('description'), location=request.POST.get('location'), area=request.POST.get('area'), price=request.POST.get('price'), is_active=request.POST.get('is_active') == 'on', is_featured=request.POST.get('is_featured') == 'on')
        if request.FILES.get('main_image'):
            p.main_image = request.FILES['main_image']
        p.save()
        for img in request.FILES.getlist('gallery_images'):
            PropertyImage.objects.create(property=p, image=img)
        messages.success(request, "Terrain créé.")
        return redirect('admin_properties')
    return render(request, 'adminpanel/properties/form.html', {'action': 'Créer'})


@admin_required
def admin_property_edit(request, pk):
    p = get_object_or_404(Property, pk=pk)
    if request.method == 'POST':
        p.title = request.POST.get('title')
        p.description = request.POST.get('description')
        p.location = request.POST.get('location')
        p.area = request.POST.get('area')
        p.price = request.POST.get('price')
        p.is_active = request.POST.get('is_active') == 'on'
        p.is_featured = request.POST.get('is_featured') == 'on'
        if request.FILES.get('main_image'):
            p.main_image = request.FILES['main_image']
        p.save()
        for img in request.FILES.getlist('gallery_images'):
            PropertyImage.objects.create(property=p, image=img)
        messages.success(request, "Terrain modifié.")
        return redirect('admin_properties')
    return render(request, 'adminpanel/properties/form.html', {'property': p, 'action': 'Modifier'})


@admin_required
def admin_property_delete(request, pk):
    p = get_object_or_404(Property, pk=pk)
    if request.method == 'POST':
        p.delete()
        messages.success(request, "Terrain supprimé.")
    return redirect('admin_properties')


# ── PROMOTIONS ──
@admin_required
def admin_promotions(request):
    return render(request, 'adminpanel/promotions/list.html', {'promotions': Promotion.objects.all()})


@admin_required
def admin_promotion_create(request):
    if request.method == 'POST':
        pr = Promotion(title=request.POST.get('title'), subtitle=request.POST.get('subtitle', ''), description=request.POST.get('description'), badge=request.POST.get('badge', ''), badge_color=request.POST.get('badge_color', '#C1121F'), price=request.POST.get('price') or None, location=request.POST.get('location', ''), advantages=request.POST.get('advantages', ''), conditions=request.POST.get('conditions', ''), order=int(request.POST.get('order', 0)), is_active=request.POST.get('is_active') == 'on')
        if request.FILES.get('image'):
            pr.image = request.FILES['image']
        pr.save()
        messages.success(request, "Promotion créée.")
        return redirect('admin_promotions')
    return render(request, 'adminpanel/promotions/form.html', {'action': 'Créer'})


@admin_required
def admin_promotion_edit(request, pk):
    pr = get_object_or_404(Promotion, pk=pk)
    if request.method == 'POST':
        pr.title = request.POST.get('title')
        pr.subtitle = request.POST.get('subtitle', '')
        pr.description = request.POST.get('description')
        pr.badge = request.POST.get('badge', '')
        pr.badge_color = request.POST.get('badge_color', '#C1121F')
        pr.price = request.POST.get('price') or None
        pr.location = request.POST.get('location', '')
        pr.advantages = request.POST.get('advantages', '')
        pr.conditions = request.POST.get('conditions', '')
        pr.order = int(request.POST.get('order', 0))
        pr.is_active = request.POST.get('is_active') == 'on'
        if request.FILES.get('image'):
            pr.image = request.FILES['image']
        pr.save()
        messages.success(request, "Promotion modifiée.")
        return redirect('admin_promotions')
    return render(request, 'adminpanel/promotions/form.html', {'promotion': pr, 'action': 'Modifier'})


@admin_required
def admin_promotion_delete(request, pk):
    pr = get_object_or_404(Promotion, pk=pk)
    if request.method == 'POST':
        pr.delete()
        messages.success(request, "Promotion supprimée.")
    return redirect('admin_promotions')


# ── INTITULÉS ──
@admin_required
def admin_intitules(request):
    return render(request, 'adminpanel/intitules/list.html', {'intitules': IntituleDossier.objects.all()})


@admin_required
def admin_intitule_create(request):
    if request.method == 'POST':
        obj = IntituleDossier(
            name=request.POST.get('name'),
            description=request.POST.get('description', '')
        )
        if request.FILES.get('image'):
            obj.image = request.FILES['image']
        obj.save()
        messages.success(request, "Intitulé créé.")
        return redirect('admin_intitules')
    return render(request, 'adminpanel/intitules/form.html', {'action': 'Créer'})


@admin_required
def admin_intitule_edit(request, pk):
    obj = get_object_or_404(IntituleDossier, pk=pk)
    if request.method == 'POST':
        obj.name = request.POST.get('name')
        obj.description = request.POST.get('description', '')
        if request.FILES.get('image'):
            obj.image = request.FILES['image']
        obj.save()
        messages.success(request, "Intitulé modifié.")
        return redirect('admin_intitules')
    return render(request, 'adminpanel/intitules/form.html', {'intitule': obj, 'action': 'Modifier'})


@admin_required
def admin_intitule_delete(request, pk):
    obj = get_object_or_404(IntituleDossier, pk=pk)
    if request.method == 'POST':
        obj.delete()
        messages.success(request, "Intitulé supprimé.")
    return redirect('admin_intitules')


# ── UTILISATEURS ──
@admin_required
def admin_clients(request):
    role_filter = request.GET.get('role', '')
    sans_dossier = request.GET.get('sans_dossier', '')
    date_debut = request.GET.get('date_debut', '')
    date_fin = request.GET.get('date_fin', '')

    users = User.objects.all().order_by('-date_joined')
    if role_filter:
        users = users.filter(role=role_filter)
    if sans_dossier == '1':
        users = users.filter(dossiers__isnull=True)
    if date_debut:
        try:
            users = users.filter(date_joined__date__gte=datetime.datetime.strptime(date_debut, '%Y-%m-%d').date())
        except ValueError:
            pass
    if date_fin:
        try:
            users = users.filter(date_joined__date__lte=datetime.datetime.strptime(date_fin, '%Y-%m-%d').date())
        except ValueError:
            pass

    return render(request, 'adminpanel/clients/list.html', {
        'clients': users,
        'role_filter': role_filter,
        'roles': User.ROLE_CHOICES,
        'sans_dossier': sans_dossier,
        'date_debut': date_debut,
        'date_fin': date_fin,
    })

@admin_required
def admin_client_create(request):
    if request.method == 'POST':
        username = request.POST.get('username')
        password = request.POST.get('password')
        role = request.POST.get('role', 'client')
        if User.objects.filter(username=username).exists():
            messages.error(request, "Cet identifiant existe déjà.")
            return render(request, 'adminpanel/clients/form.html', {'action': 'Créer', 'roles': User.ROLE_CHOICES})
        u = User.objects.create_user(
            username=username, password=password,
            last_name=request.POST.get('last_name', ''),
            email=request.POST.get('email', ''),
            phone=request.POST.get('phone', ''),
            city=request.POST.get('city', ''),
            role=role,
            sexe=request.POST.get('sexe', 'masculin'),
        )
        if role in ('admin', 'commercial'):
            u.is_staff = True
            u.save()
        messages.success(request, f"Utilisateur créé — Identifiant : {username} | Mot de passe : {password}")
        return redirect('admin_clients')
    return render(request, 'adminpanel/clients/form.html', {'action': 'Créer', 'roles': User.ROLE_CHOICES})


@admin_required
def admin_client_edit(request, pk):
    u = get_object_or_404(User, pk=pk)
    if request.method == 'POST':
        new_username = request.POST.get('username', '').strip()
        if new_username and new_username != u.username:
            if User.objects.filter(username=new_username).exclude(pk=pk).exists():
                messages.error(request, f"L'identifiant '{new_username}' est déjà utilisé.")
                return render(request, 'adminpanel/clients/form.html', {'client': u, 'action': 'Modifier', 'roles': User.ROLE_CHOICES})
            u.username = new_username
        u.last_name = request.POST.get('last_name', '')
        u.email = request.POST.get('email', '')
        u.phone = request.POST.get('phone', '')
        u.city = request.POST.get('city', '')
        u.is_active = request.POST.get('is_active') == 'on'
        u.sexe = request.POST.get('sexe', 'masculin')
        new_role = request.POST.get('role', u.role)
        u.role = new_role
        u.is_staff = new_role in ('admin', 'commercial')
        new_password = request.POST.get('password', '').strip()
        if new_password:
            u.set_password(new_password)
        u.save()
        messages.success(request, f"Utilisateur '{u.username}' modifié avec succès.")
        return redirect('admin_clients')
    return render(request, 'adminpanel/clients/form.html', {
        'client': u,
        'action': 'Modifier',
        'roles': User.ROLE_CHOICES
    })

@admin_required
def admin_client_delete(request, pk):
    u = get_object_or_404(User, pk=pk)
    if request.method == 'POST':
        u.delete()
        messages.success(request, "Utilisateur supprimé.")
    return redirect('admin_clients')


@admin_required
def admin_clients_export(request):
    wb = Workbook()
    ws = wb.active
    ws.title = "Clients et Dossiers"

    header_fill = PatternFill("solid", fgColor="0B1F3A")
    header_font = Font(bold=True, color="FFFFFF")
    headers = [
        'Identifiant*', 'Mot de passe*', 'Nom', 'Email',
        'Téléphone', 'Ville', 'Sexe', 'Rôle', 'Actif',
        'Intitulé dossier*', 'Superficie (m²)', 'Date paiement (YYYY-MM-DD)', 'Notes dossier'
    ]
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 22

    clients = User.objects.filter(role='client').order_by('last_name').prefetch_related('dossiers__intitule')
    for u in clients:
        dossiers = list(u.dossiers.all())
        if dossiers:
            for d in dossiers:
                ws.append([
                    u.username,
                    '(non exporté)',
                    u.last_name,
                    u.email,
                    u.phone,
                    u.city,
                    u.sexe,
                    u.role,
                    'Oui' if u.is_active else 'Non',
                    d.intitule.name if d.intitule else '',
                    float(d.superficie) if d.superficie else '',
                    d.date_paiement.strftime('%Y-%m-%d') if d.date_paiement else '',
                    d.description or '',
                ])
        else:
            ws.append([
                u.username, '(non exporté)', u.last_name,
                u.email, u.phone, u.city, u.sexe, u.role,
                'Oui' if u.is_active else 'Non',
                '', '', '', ''
            ])

    col_widths = [18, 18, 20, 24, 18, 16, 12, 14, 8, 22, 14, 22, 24]
    for i, w in enumerate(col_widths, 1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w

    ws2 = wb.create_sheet("Instructions")
    ws2['A1'] = "INSTRUCTIONS D'IMPORT / EXPORT"
    ws2['A1'].font = Font(bold=True, size=13, color="C1121F")
    lignes = [
        ("Une seule feuille", "Toutes les données clients ET dossiers sont sur la feuille 'Clients et Dossiers'"),
        ("Colonnes 1-9", "Informations du client : Identifiant*, Mot de passe*, Nom, Email, Téléphone, Ville, Sexe, Rôle, Actif"),
        ("Colonnes 10-13", "Informations du dossier : Intitulé*, Superficie, Date paiement, Notes"),
        ("Client existant", "Si l'identifiant existe déjà en base → le dossier est créé et lié au client existant"),
        ("Nouveau client", "Si l'identifiant est nouveau → le client ET le dossier sont créés"),
        ("Plusieurs dossiers", "Répétez l'identifiant du client sur plusieurs lignes avec des intitulés différents"),
        ("Pas de dossier", "Laissez la colonne 'Intitulé dossier' vide pour créer un client sans dossier"),
        ("Sexe", "masculin / feminin / plusieurs"),
        ("Rôle", "client / commercial / admin"),
        ("Date", "Format YYYY-MM-DD (ex: 2024-06-15) ou JJ/MM/AAAA"),
        ("* = obligatoire", "Identifiant toujours obligatoire. Mot de passe obligatoire seulement si nouveau client."),
        ("Intitulés existants", ", ".join(IntituleDossier.objects.values_list('name', flat=True)) or "Aucun — sera créé automatiquement"),
    ]
    for i, (t, d) in enumerate(lignes, 3):
        ws2.cell(row=i, column=1, value=t).font = Font(bold=True)
        ws2.cell(row=i, column=2, value=d)
    ws2.column_dimensions['A'].width = 22
    ws2.column_dimensions['B'].width = 72

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename="eden_clients_dossiers.xlsx"'
    wb.save(response)
    return response


@admin_required
def admin_clients_import(request):
    if request.method == 'POST' and request.FILES.get('excel_file'):
        excel_file = request.FILES['excel_file']
        try:
            wb = openpyxl.load_workbook(excel_file, data_only=True)
        except Exception as e:
            messages.error(request, f"Fichier invalide : {e}")
            return redirect('admin_clients_import')

        sheet_name = None
        for candidate in ['Clients et Dossiers', 'Clients & Dossiers', 'Sheet1', 'Feuil1']:
            if candidate in wb.sheetnames:
                sheet_name = candidate
                break
        if not sheet_name:
            sheet_name = wb.sheetnames[0]

        ws = wb[sheet_name]
        rows = list(ws.iter_rows(min_row=2))

        created_users = 0
        updated_users = 0
        created_dossiers = 0
        errors = []
        seen_clients = {}

        for i, row in enumerate(rows, 2):
            def get(idx):
                return _xl_val(row[idx] if len(row) > idx else None)

            username = get(0)
            if not username:
                continue

            password = get(1)
            nom = get(2)
            email = get(3)
            phone = get(4)
            city = get(5)
            sexe_raw = get(6).lower()
            role_raw = get(7).lower()
            actif_raw = get(8).lower()
            intitule_name = get(9)
            superficie_raw = get(10)
            date_raw = get(11)
            description = get(12)

            sexe_map = {'masculin': 'masculin', 'féminin': 'feminin', 'feminin': 'feminin', 'plusieurs': 'plusieurs'}
            sexe = sexe_map.get(sexe_raw, 'masculin')
            role_map = {'admin': 'admin', 'administrateur': 'admin', 'commercial': 'commercial', 'client': 'client'}
            role = role_map.get(role_raw, 'client')
            is_active = actif_raw not in ('non', 'false', '0', 'no')

            # ── Gestion client ──
            if username in seen_clients:
                user = seen_clients[username]
            else:
                try:
                    user = User.objects.get(username=username)
                    seen_clients[username] = user
                    updated_users += 1
                except User.DoesNotExist:
                    if not password:
                        errors.append(f"Ligne {i} : mot de passe manquant pour le nouveau client '{username}'")
                        continue
                    try:
                        user = User.objects.create_user(
                            username=username,
                            password=password,
                            last_name=nom,
                            email=email,
                            phone=phone,
                            city=city,
                            role=role,
                            sexe=sexe,
                            is_active=is_active,
                        )
                        if role in ('admin', 'commercial'):
                            user.is_staff = True
                            user.save()
                        seen_clients[username] = user
                        created_users += 1
                    except Exception as e:
                        errors.append(f"Ligne {i} : erreur création client '{username}' — {e}")
                        continue

            # ── Gestion dossier ──
            if not intitule_name:
                continue

            intitule = IntituleDossier.objects.filter(name__iexact=intitule_name).first()
            if not intitule:
                intitule = IntituleDossier.objects.create(name=intitule_name)

            try:
                superficie = float(superficie_raw) if superficie_raw else None
            except ValueError:
                superficie = None

            date_paiement = None
            if date_raw:
                for fmt_str in ('%Y-%m-%d', '%d/%m/%Y', '%d-%m-%Y', '%Y/%m/%d'):
                    try:
                        date_paiement = datetime.datetime.strptime(date_raw, fmt_str).date()
                        break
                    except ValueError:
                        continue

            try:
                d = Dossier(
                    client=user,
                    intitule=intitule,
                    superficie=superficie,
                    date_paiement=date_paiement,
                    description=description,
                )
                d.save()
                d.planifier_cochages_automatiques()
                DossierHistorique.objects.create(
                    dossier=d,
                    message="Dossier créé par import Excel"
                )
                created_dossiers += 1
            except Exception as e:
                errors.append(f"Ligne {i} : erreur création dossier '{intitule_name}' pour '{username}' — {e}")

        msg = (
            f"Import terminé : {created_users} client(s) créé(s), "
            f"{updated_users} client(s) existant(s) reconnu(s), "
            f"{created_dossiers} dossier(s) créé(s)."
        )
        if errors:
            msg += f" {len(errors)} avertissement(s)."
            for err in errors[:10]:
                messages.warning(request, err)
        messages.success(request, msg)
        return redirect('admin_clients')

    return render(request, 'adminpanel/clients/import_excel.html')


@admin_required
def admin_clients_template(request):
    wb = Workbook()
    ws = wb.active
    ws.title = "Clients et Dossiers"

    header_fill = PatternFill("solid", fgColor="0B1F3A")
    header_font = Font(bold=True, color="FFFFFF")
    red_fill = PatternFill("solid", fgColor="FFF0F0")

    headers = [
        'Identifiant*', 'Mot de passe*', 'Nom', 'Email',
        'Téléphone', 'Ville', 'Sexe', 'Rôle', 'Actif',
        'Intitulé dossier', 'Superficie (m²)', 'Date paiement (YYYY-MM-DD)', 'Notes dossier'
    ]
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 22

    examples = [
        ['jean.mballa', 'pass123', 'Mballa Jean', 'jean@mail.cm', '+237600000001', 'Yaoundé', 'masculin', 'client', 'Oui', 'Lotissement Odza', 600, '2024-03-15', 'Parcelle A12'],
        ['jean.mballa', '', '', '', '', '', '', '', '', 'Parcelle Bastos', 400, '2024-04-01', 'Lot B5'],
        ['marie.ngo', 'pass456', 'Ngo Marie', 'marie@mail.cm', '+237600000002', 'Douala', 'feminin', 'client', 'Oui', 'Nkolbisson', 800, '2024-05-10', ''],
        ['paul.biya', 'pass789', 'Biya Paul', '', '+237600000003', 'Yaoundé', 'masculin', 'client', 'Oui', '', '', '', ''],
    ]
    note_fill = PatternFill("solid", fgColor="E8F0FE")
    for row_idx, ex in enumerate(examples, 2):
        for col_idx, val in enumerate(ex, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            if row_idx == 3:
                cell.fill = note_fill

    col_widths = [18, 18, 20, 24, 18, 16, 12, 12, 8, 22, 14, 24, 24]
    for i, w in enumerate(col_widths, 1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w

    ws2 = wb.create_sheet("Instructions")
    ws2['A1'] = "GUIDE D'UTILISATION"
    ws2['A1'].font = Font(bold=True, size=13, color="C1121F")
    lignes = [
        ("RÈGLE PRINCIPALE", "Une seule feuille pour tout : clients ET leurs dossiers"),
        ("Client existant", "Si l'identifiant existe déjà → le dossier est ajouté au client existant (mot de passe ignoré)"),
        ("Nouveau client", "Si l'identifiant est nouveau → client + dossier créés (mot de passe obligatoire)"),
        ("Plusieurs dossiers", "Ligne 2 exemple : jean.mballa a 2 dossiers → répéter l'identifiant, laisser les infos client vides"),
        ("Client sans dossier", "Ligne 4 exemple : paul.biya est créé sans dossier (colonnes 10-13 vides)"),
        ("Sexe", "masculin / feminin / plusieurs"),
        ("Rôle", "client / commercial / admin (défaut : client)"),
        ("Date paiement", "Format YYYY-MM-DD (2024-06-15) ou JJ/MM/AAAA (15/06/2024)"),
        ("Intitulé", "S'il n'existe pas en base, il sera créé automatiquement"),
        ("Intitulés existants", ", ".join(IntituleDossier.objects.values_list('name', flat=True)) or "Aucun pour le moment"),
        ("* Obligatoire", "Identifiant toujours requis. Mot de passe requis seulement pour les nouveaux clients."),
    ]
    for i, (t, d) in enumerate(lignes, 3):
        ws2.cell(row=i, column=1, value=t).font = Font(bold=True)
        ws2.cell(row=i, column=2, value=d)
    ws2.column_dimensions['A'].width = 22
    ws2.column_dimensions['B'].width = 72

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename="modele_import_eden.xlsx"'
    wb.save(response)
    return response


# ── ÉTAPES ──
@admin_required
def admin_etapes_globales(request):
    etapes_tech = EtapeGlobale.objects.filter(type='technique')
    etapes_morc = EtapeGlobale.objects.filter(type='morcellement')
    return render(request, 'adminpanel/etapes/list.html', {'etapes_tech': etapes_tech, 'etapes_morc': etapes_morc, 'total_tech': sum(e.percentage for e in etapes_tech), 'total_morc': sum(e.percentage for e in etapes_morc)})


@admin_required
def admin_etape_globale_create(request):
    if request.method == 'POST':
        type_ = request.POST.get('type')
        pct = int(request.POST.get('percentage', 0))
        total_actuel = sum(e.percentage for e in EtapeGlobale.objects.filter(type=type_))
        if total_actuel + pct > 100:
            messages.error(request, f"Impossible : total dépasserait 100% (actuel : {total_actuel}%)")
            return redirect('admin_etapes_globales')
        EtapeGlobale.objects.create(name=request.POST.get('name'), type=type_, percentage=pct, order=int(request.POST.get('order', 0)), description=request.POST.get('description', ''))
        messages.success(request, "Étape créée.")
        return redirect('admin_etapes_globales')
    return render(request, 'adminpanel/etapes/form.html', {'action': 'Créer'})


@admin_required
def admin_etape_globale_edit(request, pk):
    etape = get_object_or_404(EtapeGlobale, pk=pk)
    if request.method == 'POST':
        pct = int(request.POST.get('percentage', 0))
        total_actuel = sum(e.percentage for e in EtapeGlobale.objects.filter(type=etape.type).exclude(pk=pk))
        if total_actuel + pct > 100:
            messages.error(request, "Impossible : total dépasserait 100%")
            return redirect('admin_etapes_globales')
        etape.name = request.POST.get('name')
        etape.percentage = pct
        etape.order = int(request.POST.get('order', 0))
        etape.description = request.POST.get('description', '')
        etape.save()
        messages.success(request, "Étape modifiée.")
        return redirect('admin_etapes_globales')
    return render(request, 'adminpanel/etapes/form.html', {'etape': etape, 'action': 'Modifier'})


@admin_required
def admin_etape_globale_delete(request, pk):
    etape = get_object_or_404(EtapeGlobale, pk=pk)
    if request.method == 'POST':
        etape.delete()
        messages.success(request, "Étape supprimée.")
    return redirect('admin_etapes_globales')


# ── DOSSIERS ──
@admin_required
def admin_dossiers(request):
    date_debut = request.GET.get('date_debut', '')
    date_fin = request.GET.get('date_fin', '')
    sans_superficie = request.GET.get('sans_superficie', '')

    dossiers = Dossier.objects.select_related('client', 'intitule').all()

    if date_debut:
        try:
            dossiers = dossiers.filter(created_at__date__gte=datetime.datetime.strptime(date_debut, '%Y-%m-%d').date())
        except ValueError:
            pass
    if date_fin:
        try:
            dossiers = dossiers.filter(created_at__date__lte=datetime.datetime.strptime(date_fin, '%Y-%m-%d').date())
        except ValueError:
            pass
    if sans_superficie == '1':
        dossiers = dossiers.filter(superficie__isnull=True)

    dossiers_data = [{'dossier': d, 'is_complete': d.is_complete()} for d in dossiers]
    etapes_tech = EtapeGlobale.objects.filter(type='technique')
    etapes_morc = EtapeGlobale.objects.filter(type='morcellement')
    return render(request, 'adminpanel/dossiers/list.html', {
        'dossiers_data': dossiers_data,
        'etapes_tech': etapes_tech,
        'etapes_morc': etapes_morc,
        'date_debut': date_debut,
        'date_fin': date_fin,
        'sans_superficie': sans_superficie,
    })

@admin_required
def admin_dossier_create(request):
    clients = User.objects.filter(role='client', is_active=True)
    intitules = IntituleDossier.objects.all()
    if request.method == 'POST':
        client = get_object_or_404(User, pk=request.POST.get('client'), role='client')
        intitule_id = request.POST.get('intitule')
        intitule = get_object_or_404(IntituleDossier, pk=intitule_id) if intitule_id else None
        d = Dossier(client=client, intitule=intitule, description=request.POST.get('description', ''), superficie=request.POST.get('superficie') or None, date_paiement=request.POST.get('date_paiement') or None)
        if request.FILES.get('photo'):
            d.photo = request.FILES['photo']
        d.save()
        d.planifier_cochages_automatiques()
        DossierHistorique.objects.create(dossier=d, message="Dossier créé — cochage automatique planifié", created_by=request.user)
        messages.success(request, f"Dossier {d.reference} créé.")
        return redirect('admin_dossier_detail', pk=d.pk)
    return render(request, 'adminpanel/dossiers/form.html', {'clients': clients, 'intitules': intitules, 'action': 'Créer'})


@admin_required
def admin_dossier_detail(request, pk):
    dossier = get_object_or_404(Dossier, pk=pk)
    dossier.appliquer_cochages_automatiques()
    etapes_tech = dossier.get_etapes_technique()
    etapes_morc = dossier.get_etapes_morcellement()
    historique = dossier.historique.all()[:30]
    prog_tech = dossier.get_progression_technique()
    prog_morc = dossier.get_progression_morcellement()
    total_tech = dossier.get_total_technique()
    total_morc = dossier.get_total_morcellement()
    pct_tech_bar = round(prog_tech / total_tech * 100) if total_tech > 0 else 0
    pct_morc_bar = round(prog_morc / total_morc * 100) if total_morc > 0 else 0
    etapes_tech_list = list(EtapeGlobale.objects.filter(type='technique'))
    etapes_morc_list = list(EtapeGlobale.objects.filter(type='morcellement'))
    return render(request, 'adminpanel/dossiers/detail.html', {
        'dossier': dossier, 'etapes_tech': etapes_tech, 'etapes_morc': etapes_morc,
        'historique': historique, 'prog_tech': prog_tech, 'prog_morc': prog_morc,
        'total_tech': total_tech, 'total_morc': total_morc,
        'pct_tech_bar': pct_tech_bar, 'pct_morc_bar': pct_morc_bar,
        'current_tech': dossier.get_current_etape_technique(),
        'current_morc': dossier.get_current_etape_morcellement(),
        'morc_locked': not dossier.is_technique_complete(),
        'last_tech_id': etapes_tech_list[-1].pk if etapes_tech_list else None,
        'last_morc_id': etapes_morc_list[-1].pk if etapes_morc_list else None,
    })


@admin_required
@require_POST
def admin_etape_toggle(request, dossier_pk, etape_pk):
    dossier = get_object_or_404(Dossier, pk=dossier_pk)
    etape = get_object_or_404(EtapeGlobale, pk=etape_pk)
    if etape.type == 'morcellement' and not dossier.is_technique_complete():
        return JsonResponse({'status': 'error', 'message': 'Le technique doit être à 100% d\'abord.'}, status=403)
    cochee, _ = EtapeCochee.objects.get_or_create(dossier=dossier, etape=etape)
    cochee.is_done = not cochee.is_done
    cochee.done_at = timezone.now() if cochee.is_done else None
    cochee.done_by = request.user if cochee.is_done else None
    if not cochee.is_done:
        cochee.done_by = None
    cochee.save()
    action = "validée manuellement" if cochee.is_done else "annulée"
    prog = dossier.get_progression_technique() if etape.type == 'technique' else dossier.get_progression_morcellement()
    total = dossier.get_total_technique() if etape.type == 'technique' else dossier.get_total_morcellement()
    DossierHistorique.objects.create(dossier=dossier, message=f"Étape '{etape.name}' {action} — {etape.get_type_display()} : {prog}%", created_by=request.user)
    pct_bar = round(prog / total * 100) if total > 0 else 0
    current = dossier.get_current_etape_technique() if etape.type == 'technique' else dossier.get_current_etape_morcellement()
    return JsonResponse({'status': 'ok', 'is_done': cochee.is_done, 'progression': prog, 'pct_bar': pct_bar, 'type': etape.type, 'morc_unlocked': dossier.is_technique_complete(), 'current_etape_name': current.name if current else '', 'current_etape_desc': current.description if current else ''})


@admin_required
def admin_dossier_edit(request, pk):
    dossier = get_object_or_404(Dossier, pk=pk)
    clients = User.objects.filter(role='client', is_active=True)
    intitules = IntituleDossier.objects.all()
    if request.method == 'POST':
        client_id = request.POST.get('client')
        if client_id:
            dossier.client_id = client_id
        intitule_id = request.POST.get('intitule')
        dossier.intitule = get_object_or_404(IntituleDossier, pk=intitule_id) if intitule_id else None
        dossier.superficie = request.POST.get('superficie') or None
        dossier.date_paiement = request.POST.get('date_paiement') or None
        dossier.description = request.POST.get('description', '')
        if request.FILES.get('photo'):
            dossier.photo = request.FILES['photo']
        dossier.save()
        DossierHistorique.objects.create(
            dossier=dossier,
            message="Dossier modifié",
            created_by=request.user
        )
        messages.success(request, "Dossier modifié avec succès.")
        return redirect('admin_dossier_detail', pk=dossier.pk)
    return render(request, 'adminpanel/dossiers/edit.html', {
        'dossier': dossier,
        'clients': clients,
        'intitules': intitules,
    })


@admin_required
def admin_dossier_delete(request, pk):
    d = get_object_or_404(Dossier, pk=pk)
    if request.method == 'POST':
        d.delete()
        messages.success(request, "Dossier supprimé.")
        return redirect('admin_dossiers')
    return render(request, 'adminpanel/dossiers/confirm_delete.html', {'dossier': d})


@admin_required
def admin_dossiers_alerte(request):
    import datetime as dt
    today = timezone.now().date()
    seuil = today - dt.timedelta(days=45)
    tous = Dossier.objects.filter(
        date_paiement__isnull=False,
        date_paiement__lte=seuil
    ).select_related('client', 'intitule')
    dossiers = [d for d in tous if not d.is_complete()]

    if request.GET.get('export') == 'csv':
        response = HttpResponse(content_type='text/csv; charset=utf-8')
        response['Content-Disposition'] = 'attachment; filename="dossiers_alerte_45j.csv"'
        response.write('\ufeff')
        writer = csv.writer(response, delimiter=';')
        writer.writerow(['Référence', 'Intitulé', 'Client', 'Téléphone', 'Email', 'Superficie', 'Date paiement', 'Jours depuis paiement', 'Tech %', 'Morc %'])
        for d in dossiers:
            jours = (today - d.date_paiement).days if d.date_paiement else 0
            writer.writerow([
                d.reference, d.get_title(),
                d.client.get_full_name() or d.client.username,
                d.client.phone, d.client.email,
                d.superficie or '',
                d.date_paiement.strftime('%d/%m/%Y') if d.date_paiement else '',
                jours,
                d.get_progression_technique(),
                d.get_progression_morcellement()
            ])
        return response

    dossiers_with_days = []
    for d in dossiers:
        jours = (today - d.date_paiement).days if d.date_paiement else 0
        dossiers_with_days.append({'dossier': d, 'jours': jours})

    return render(request, 'adminpanel/dossiers/alerte.html', {
        'dossiers_with_days': dossiers_with_days,
        'count': len(dossiers_with_days),
    })

# ── RÉSERVATIONS ──
@admin_required
def admin_reservations(request):
    return render(request, 'adminpanel/reservations/list.html', {'reservations': Reservation.objects.select_related('property_link', 'promotion_link').all()})


@admin_required
def admin_reservation_update(request, pk):
    r = get_object_or_404(Reservation, pk=pk)
    if request.method == 'POST':
        r.status = request.POST.get('status', r.status)
        r.save()
        messages.success(request, "Statut mis à jour.")
    return redirect('admin_reservations')


# ── CHATBOT ──
@admin_required
def admin_chatbot(request):
    return render(request, 'adminpanel/chatbot/list.html', {'knowledges': ChatbotKnowledge.objects.all(), 'unknowns': ChatbotUnknown.objects.all()[:50]})


@admin_required
def admin_chatbot_create(request):
    if request.method == 'POST':
        ChatbotKnowledge.objects.create(keywords=request.POST.get('keywords'), question_example=request.POST.get('question_example'), answer=request.POST.get('answer'), order=int(request.POST.get('order', 0)), is_active=request.POST.get('is_active') == 'on')
        messages.success(request, "Réponse ajoutée.")
        return redirect('admin_chatbot')
    return render(request, 'adminpanel/chatbot/form.html', {'action': 'Ajouter'})


@admin_required
def admin_chatbot_edit(request, pk):
    kb = get_object_or_404(ChatbotKnowledge, pk=pk)
    if request.method == 'POST':
        kb.keywords = request.POST.get('keywords')
        kb.question_example = request.POST.get('question_example')
        kb.answer = request.POST.get('answer')
        kb.order = int(request.POST.get('order', 0))
        kb.is_active = request.POST.get('is_active') == 'on'
        kb.save()
        messages.success(request, "Réponse modifiée.")
        return redirect('admin_chatbot')
    return render(request, 'adminpanel/chatbot/form.html', {'kb': kb, 'action': 'Modifier'})


@admin_required
def admin_chatbot_delete(request, pk):
    kb = get_object_or_404(ChatbotKnowledge, pk=pk)
    if request.method == 'POST':
        kb.delete()
        messages.success(request, "Supprimé.")
    return redirect('admin_chatbot')

@admin_required
def admin_clients_bulk(request):
    if request.method == 'POST':
        ids = request.POST.getlist('selected_ids')
        action = request.POST.get('bulk_action')
        if not ids:
            messages.error(request, "Aucun utilisateur sélectionné.")
            return redirect('admin_clients')
        users = User.objects.filter(pk__in=ids)
        if action == 'activer':
            users.update(is_active=True)
            messages.success(request, f"{users.count()} utilisateur(s) activé(s).")
        elif action == 'desactiver':
            users.exclude(pk=request.user.pk).update(is_active=False)
            messages.success(request, "Utilisateurs désactivés.")
        elif action == 'sexe_masculin':
            users.update(sexe='masculin')
            messages.success(request, "Sexe mis à jour : Masculin.")
        elif action == 'sexe_feminin':
            users.update(sexe='feminin')
            messages.success(request, "Sexe mis à jour : Féminin.")
        elif action == 'sexe_plusieurs':
            users.update(sexe='plusieurs')
            messages.success(request, "Sexe mis à jour : Plusieurs.")
        elif action == 'supprimer':
            count = users.exclude(pk=request.user.pk).count()
            users.exclude(pk=request.user.pk).delete()
            messages.success(request, f"{count} utilisateur(s) supprimé(s).")
    return redirect('admin_clients')


@admin_required
def admin_dossiers_bulk(request):
    if request.method == 'POST':
        ids = request.POST.getlist('selected_ids')
        action = request.POST.get('bulk_action')
        if not ids:
            messages.error(request, "Aucun dossier sélectionné.")
            return redirect('admin_dossiers')
        dossiers = Dossier.objects.filter(pk__in=ids)

        if action == 'supprimer':
            count = dossiers.count()
            dossiers.delete()
            messages.success(request, f"{count} dossier(s) supprimé(s).")

        elif action == 'export_excel':
            return _export_dossiers_excel(dossiers)

        elif action == 'avancement_tech':
            etape_ids = request.POST.getlist('etapes_tech_ids')
            for dossier in dossiers:
                for etape in EtapeGlobale.objects.filter(type='technique'):
                    cochee, _ = EtapeCochee.objects.get_or_create(dossier=dossier, etape=etape)
                    should_done = str(etape.pk) in etape_ids
                    if cochee.is_done != should_done:
                        cochee.is_done = should_done
                        cochee.done_at = timezone.now() if should_done else None
                        cochee.done_by = request.user if should_done else None
                        cochee.save()
                DossierHistorique.objects.create(
                    dossier=dossier,
                    message=f"Avancement technique mis à jour en masse ({dossier.get_progression_technique()}%)",
                    created_by=request.user
                )
            messages.success(request, f"Avancement technique appliqué à {dossiers.count()} dossier(s).")

        elif action == 'avancement_morc':
            if not all(d.is_technique_complete() for d in dossiers):
                messages.error(request, "Certains dossiers sélectionnés n'ont pas leur technique à 100%. Impossible d'appliquer le morcellement.")
                return redirect('admin_dossiers')
            etape_ids = request.POST.getlist('etapes_morc_ids')
            for dossier in dossiers:
                for etape in EtapeGlobale.objects.filter(type='morcellement'):
                    cochee, _ = EtapeCochee.objects.get_or_create(dossier=dossier, etape=etape)
                    should_done = str(etape.pk) in etape_ids
                    if cochee.is_done != should_done:
                        cochee.is_done = should_done
                        cochee.done_at = timezone.now() if should_done else None
                        cochee.done_by = request.user if should_done else None
                        cochee.save()
                DossierHistorique.objects.create(
                    dossier=dossier,
                    message=f"Avancement morcellement mis à jour en masse ({dossier.get_progression_morcellement()}%)",
                    created_by=request.user
                )
            messages.success(request, f"Avancement morcellement appliqué à {dossiers.count()} dossier(s).")

        elif action == 'superficie':
            nouvelle_superficie = request.POST.get('nouvelle_superficie', '').strip()
            if not nouvelle_superficie:
                messages.error(request, "Veuillez saisir une superficie.")
                return redirect('admin_dossiers')
            try:
                val = float(nouvelle_superficie)
                count = dossiers.count()
                dossiers.update(superficie=val)
                for d in Dossier.objects.filter(pk__in=ids):
                        DossierHistorique.objects.create(
                        dossier=d,
                        message=f"Superficie mise à jour en masse : {val} m²",
                        created_by=request.user
                    )
                messages.success(request, f"Superficie ({val} m²) appliquée à {count} dossier(s).")
            except ValueError:
                messages.error(request, "Superficie invalide.")


    return redirect('admin_dossiers')

@admin_required
def admin_dossiers_export_excel(request):
    ids = request.GET.get('ids', '')
    status_filter = request.GET.get('status', '')
    search = request.GET.get('search', '')
    dossiers_qs = Dossier.objects.select_related('client', 'intitule').all()
    if ids:
        id_list = [i for i in ids.split(',') if i.strip().isdigit()]
        dossiers_qs = dossiers_qs.filter(pk__in=id_list)
    else:
        if status_filter == 'encours':
            dossiers_qs = [d for d in dossiers_qs if not d.is_complete()]
        elif status_filter == 'complet':
            dossiers_qs = [d for d in dossiers_qs if d.is_complete()]
        if search:
            s = search.lower()
            dossiers_qs = [d for d in dossiers_qs if s in d.reference.lower() or s in d.get_title().lower() or s in (d.client.last_name or '').lower() or s in d.client.username.lower()]
    return _export_dossiers_excel(dossiers_qs)


def _export_dossiers_excel(dossiers_qs):
    wb = Workbook()
    ws = wb.active
    ws.title = "Dossiers"
    header_fill = PatternFill("solid", fgColor="0B1F3A")
    header_font = Font(bold=True, color="FFFFFF")
    headers = ['Référence', 'Intitulé', 'Client', 'Identifiant', 'Téléphone', 'Email', 'Superficie (m²)', 'Date paiement', 'Avancement Tech (%)', 'Avancement Morc. (%)', 'Statut']
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center')
    for d in dossiers_qs:
        ws.append([
            d.reference,
            d.get_title(),
            (d.client.last_name or d.client.username).upper(),
            d.client.username,
            d.client.phone,
            d.client.email,
            float(d.superficie) if d.superficie else '',
            d.date_paiement.strftime('%d/%m/%Y') if d.date_paiement else '',
            d.get_progression_technique(),
            d.get_progression_morcellement(),
            d.get_statut_display(),
        ])
    col_widths = [18, 22, 20, 16, 16, 24, 14, 16, 16, 16, 12]
    for i, w in enumerate(col_widths, 1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename="dossiers_eden.xlsx"'
    wb.save(response)
    return response


@admin_required
def admin_dossiers_supprimer_doublons(request):
    if request.method == 'POST':
        from django.db.models import Count
        seen = {}
        doublons_ids = []
        dossiers = Dossier.objects.select_related('client', 'intitule').order_by('created_at')
        for d in dossiers:
            key = (
                d.client_id,
                d.intitule_id,
                str(d.superficie) if d.superficie else 'none'
            )
            if key in seen:
                doublons_ids.append(d.pk)
            else:
                seen[key] = d.pk
        if doublons_ids:
            count = len(doublons_ids)
            Dossier.objects.filter(pk__in=doublons_ids).delete()
            messages.success(request, f"{count} doublon(s) supprimé(s) (intitulé + client + superficie identiques — le(s) plus récent(s) supprimé(s)).")
        else:
            messages.info(request, "Aucun doublon détecté.")
    return redirect('admin_dossiers')